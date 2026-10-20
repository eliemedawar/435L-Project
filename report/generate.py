"""Generate documentation from source, schema and genuine execution records.
Developed by: Marc Bteich. This does not invent API results or imitate Postman UI.
"""
import ast
from contextlib import closing
import html
import json
from pathlib import Path
import re
import sqlite3
import subprocess
from urllib.parse import urlparse

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'report'
DIAGRAMS=REPORT/'diagrams'
EVIDENCE=REPORT/'evidence'
ATTR=' (Marc Bteich)'
SERVICES=['users','assets','incidents','threats','logs','reports','vulnerabilities','notifications','playbooks']


def diagrams(connection):
    DIAGRAMS.mkdir(exist_ok=True)
    tables=[row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
    dot=['digraph Database {','graph [rankdir=LR, bgcolor="white", pad=0.3, nodesep=0.4, ranksep=1.0, label="Database design — Marc Bteich", labelloc=t, fontname="Arial"];',
         'node [shape=plain,fontname="Arial"]; edge [fontname="Arial",fontsize=9,color="#64748b",arrowsize=0.65];']
    mermaid=['erDiagram']
    for table in tables:
        columns=connection.execute(f'PRAGMA table_info({table})').fetchall()
        foreign=connection.execute(f'PRAGMA foreign_key_list({table})').fetchall()
        fkcolumns={row[3] for row in foreign}
        rows=[]
        mermaid.append(f'    {table} {{')
        for _,name,kind,required,default,pk in columns:
            key='PK' if pk else 'FK' if name in fkcolumns else ''
            rows.append(f'<TR><TD ALIGN="LEFT">{html.escape(name)} {key}</TD><TD ALIGN="LEFT">{html.escape(kind)}</TD></TR>')
            mermaid.append(f'        {kind.split("(")[0]} {name} {key}'.rstrip())
        mermaid.append('    }')
        dot.append(f'{table} [label=<<TABLE BORDER="0" CELLBORDER="1" CELLSPACING="0" CELLPADDING="6"><TR><TD COLSPAN="2" BGCOLOR="#dbeafe"><B>{table}</B></TD></TR>'+''.join(rows)+'</TABLE>>];')
        for fk in foreign:
            dot.append(f'{fk[2]} -> {table} [label="{fk[3]}: 1 to many",dir=both,arrowtail=tee,arrowhead=crow];')
            mermaid.append(f'    {fk[2]} ||--o{{ {table} : "{fk[3]}"')
    dot.append('}')
    (DIAGRAMS/'database.dot').write_text('\n'.join(dot),encoding='utf-8')
    (DIAGRAMS/'database.mmd').write_text('\n'.join(mermaid),encoding='utf-8')
    architecture='''digraph Architecture {
      graph [rankdir=TB,bgcolor="white",fontname="Arial",label="Security backend architecture — Marc Bteich",labelloc=t,pad=0.3];
      node [shape=box,style="rounded,filled",fillcolor="#eff6ff",fontname="Arial"];
      edge [fontname="Arial",fontsize=9,color="#64748b"];
      client [label="Postman / Newman\nBearer JWT",fillcolor="#dcfce7"];
      common [label="Shared common package (Elie)\nJWT • RBAC • errors • config • DB • HTTP clients",fillcolor="#f1f5f9"];
      database [label="One shared SQLite database\n10 tables • PK/FK/CHECK/UNIQUE",shape=cylinder,fillcolor="#fef3c7"];
      users [label="1 Users (Elie)\n5001"]; assets [label="2 Assets (Marc)\n5002"];
      incidents [label="3 Incidents (Elie)\n5003"]; threats [label="4 Threats (Marc)\n5004"];
      logs [label="5 Audit Logs (Elie)\n5005"]; reports [label="6 Reports (shared)\n5006"];
      vulnerabilities [label="7 Vulnerabilities (Marc)\n5007"]; notifications [label="8 Notifications (Elie)\n5008"];
      playbooks [label="9 Playbooks (Marc)\n5009"];
      {rank=same; users;assets;incidents;} {rank=same; threats;vulnerabilities;playbooks;}
      {rank=same; logs;notifications;reports;}
      client -> common [label="authenticated public requests"];
      common -> {users assets incidents threats vulnerabilities playbooks logs notifications reports} [style=dotted,label="reuse"];
      {users assets incidents threats vulnerabilities playbooks logs notifications} -> database [style=dashed,label="ORM"];
      reports -> {users assets incidents threats vulnerabilities logs} [label="HTTP /internal/stats\nX-Service-Key",color="#2563eb",constraint=false];
      incidents -> assets [label="HTTP verify asset",constraint=false];
      vulnerabilities -> assets [label="HTTP verify asset",constraint=false];
      {assets incidents playbooks} -> users [label="HTTP verify user",constraint=false];
      {threats vulnerabilities} -> notifications [label="HTTP Critical alerts",constraint=false];
      {users assets incidents threats vulnerabilities playbooks notifications} -> logs [label="HTTP audit",constraint=false];
    }'''
    # A compact logical view avoids crossing every individual HTTP edge.
    architecture='''digraph Architecture {
      graph [rankdir=TB,bgcolor="white",fontname="Arial",label="Security backend architecture — Marc Bteich",labelloc=t,pad=0.3,nodesep=0.7,ranksep=0.65];
      node [shape=box,style="rounded,filled",fillcolor="#eff6ff",fontname="Arial",margin=0.18];
      edge [fontname="Arial",fontsize=11,color="#475569"];
      client [label="API clients: Postman / Newman\nBearer JWT on public protected endpoints",fillcolor="#dcfce7"];
      common [label="Shared Python package (Elie)\nJWT / RBAC • errors • config • SQLAlchemy • HTTP helpers",fillcolor="#f1f5f9"];
      reports [label="6 Reports — 5006\nElie: summary / incidents / activity\nMarc: assets / threats / vulnerabilities",fillcolor="#dbeafe"];
      data [shape=plain,label=<<TABLE BORDER="0" CELLBORDER="1" CELLSPACING="0" CELLPADDING="12">
        <TR><TD COLSPAN="2" BGCOLOR="#dbeafe"><B>Eight independent data-service Flask apps</B></TD></TR>
        <TR><TD>1 Users — 5001 (Elie)</TD><TD>2 Assets — 5002 (Marc)</TD></TR>
        <TR><TD>3 Incidents — 5003 (Elie)</TD><TD>4 Threats — 5004 (Marc)</TD></TR>
        <TR><TD>5 Audit Logs — 5005 (Elie)</TD><TD>7 Vulnerabilities — 5007 (Marc)</TD></TR>
        <TR><TD>8 Notifications — 5008 (Elie)</TD><TD>9 Playbooks — 5009 (Marc)</TD></TR>
      </TABLE>>];
      internal [label="Service-to-service HTTP: X-Service-Key\nVerify users / assets; write audit events; send alerts\nReports collect /internal/stats; no report table queries",fillcolor="#f8fafc"];
      database [label="One shared SQLite database\n10 tables • PK / FK / CHECK / UNIQUE\nAction execution history saved atomically",shape=cylinder,fillcolor="#fef3c7"];
      client -> data [label="REST requests"];
      client -> reports [label="admin / auditor"];
      common -> data [style=dashed,label="import/reuse",constraint=false];
      common -> reports [style=dashed,label="import/reuse",constraint=false];
      {rank=same;data;reports;}
      reports -> data [label="HTTP statistics",color="#2563eb",constraint=false];
      data -> database [label="ORM / bound values"];
      data -> internal [style=dotted,label="internal API calls"];
      {rank=same;database;internal;}
    }'''
    (DIAGRAMS/'architecture.dot').write_text(architecture,encoding='utf-8')
    (DIAGRAMS/'architecture.mmd').write_text('''flowchart TB
  Client[Postman / Newman: JWT] --> Common[Shared auth, RBAC, errors and clients]
  Common --> Services[Nine independent Flask apps: ports 5001–5009]
  Services --> DB[(Shared database: ten tables)]
  Reports[Reports] -->|HTTP service key /internal/stats| Sources[Users, Assets, Incidents, Threats, Logs, Vulnerabilities]
  B[Threats / Vulnerabilities] -->|HTTP Critical alerts| Notifications
  Services -->|HTTP write events| Audit[Audit Logs]
''',encoding='utf-8')
    node=next((ROOT/'.local').glob('node-*-win-x64/node.exe'))
    subprocess.run([str(node),str(REPORT/'render_diagrams.cjs')],check=True,cwd=ROOT)


def route_documentation():
    sections=['# 8. Merged API Documentation'+ATTR,
        'The following endpoint inventory is extracted from the actual route source. Request examples and negative tests are in the merged Postman collection. Roles shown are the route-level checks; object-level checks in Elie’s services still apply.',
        'Public authentication uses Authorization: Bearer <JWT>. Internal endpoints require X-Service-Key. JSON writes use Content-Type: application/json. Success: {status: success, message, data}; errors: {status: error, message}. Typical errors: 400 invalid input, 401 absent/invalid credentials, 403 insufficient permission, 404 missing record, 409 database conflict, 503 unavailable dependency. A 409 is only expected when an applicable constraint exists.']
    inventories=[]
    combined=json.loads((ROOT/'postman/SIDRMS.postman_collection.json').read_text())
    def walk(items):
        for item in items:
            if 'item' in item:
                yield from walk(item['item'])
            else:
                yield item
    allrequests=list(walk(combined['item']))
    executions=json.loads((EVIDENCE/'newman-combined.json').read_text())['records']
    for service in SERVICES:
        sections.append('## '+service.title()+ATTR)
        paths=[ROOT/f'services/{service}/routes.py'] if service!='reports' else [ROOT/'services/reports/routes_a.py',ROOT/'services/reports/routes_b.py']
        for path in paths:
            module=ast.parse(path.read_text())
            for function in module.body:
                if not isinstance(function,(ast.FunctionDef,ast.AsyncFunctionDef)):
                    continue
                route=None; methods=[]; roles='public'; auth='Public'
                for decorator in function.decorator_list:
                    if isinstance(decorator,ast.Call) and isinstance(decorator.func,ast.Attribute) and decorator.func.attr=='route':
                        route=ast.literal_eval(decorator.args[0])
                        methods=next(ast.literal_eval(k.value) for k in decorator.keywords if k.arg=='methods')
                    if isinstance(decorator,ast.Name) and decorator.id=='require_service':
                        roles='internal service only';auth='Service key'
                    if isinstance(decorator,ast.Call) and isinstance(decorator.func,ast.Name) and decorator.func.id=='require_auth':
                        roles=','.join((v.value if isinstance(v,ast.Constant) else v.id.lower()) for k in decorator.keywords if k.arg=='roles' for v in k.value.elts)
                        auth='Bearer JWT'
                if route is None:
                    continue
                normalized=re.sub(r'<(?:int:|string:)?([^>]+)>',r'{\1}',route)
                for method in methods:
                    sections.extend(['### '+method+' '+normalized+ATTR,
                                     f'Purpose: {function.name.replace("_"," ")}. Access: {roles}. Authentication: {auth}.'])
                    if '<' in route:
                        sections.append('Path parameters: '+', '.join(re.findall(r'<(?:int:|string:)?([^>]+)>',route))+'. IDs must identify existing records.')
                    matches=[]
                    for item in allrequests:
                        req=item['request'];url=req['url'];url=url if isinstance(url,str) else url['raw']
                        relative=re.sub(r'^\{\{base_[^}]+\}\}','',url).split('?')[0]
                        pattern=re.sub(r'<(?:int:|string:)?[^>]+>',r'[^/]+',route)
                        if req['method']==method and re.fullmatch(pattern,relative):
                            matches.append(item)
                    if matches:
                        example=next((m for m in matches if m['request'].get('body')),matches[0])
                        request=example['request'];url=request['url'];url=url if isinstance(url,str) else url['raw']
                        sections.append('Postman example: `'+method+' '+url+'`.')
                        if 'body' in request:
                            sections.extend(['Request body example (template; credentials remain private):','```json',request['body']['raw'],'```'])
                        if '/search' in route:
                            keys=set()
                            for item in matches:
                                u=item['request']['url'];u=u if isinstance(u,str) else u['raw']
                                keys.update(re.findall(r'[?&]([^=&]+)=',u))
                            sections.append('Search parameters demonstrated: '+', '.join(sorted(keys))+'; consult field rules below for allowed keys.')
                    actual_pattern=re.sub(r'<int:[^>]+>',r'[0-9]+',route)
                    actual_pattern=re.sub(r'<(?:string:)?[^>]+>',r'[^/]+',actual_pattern)
                    sample=next((record for record in executions if record['method']==method
                        and record['status'] in (200,201)
                        and re.fullmatch(actual_pattern,urlparse(record['url']).path)),None)
                    if sample:
                        response=json.dumps(sample['response'],indent=2)
                        sections.append(f'Executed response example: HTTP {sample["status"]}; Newman request "{sample["name"]}". Credentials are redacted.')
                        if len(response)<=5000:
                            sections.extend(['```json',response,'```'])
                        else:
                            sections.append('Full response is retained without truncation in evidence/newman-combined.json for that request; large list output is omitted here for readability.')
                    inventories.append(dict(service=service,method=method,path=normalized,roles=roles))
    sections.extend(['## Person B field rules and responses'+ATTR,
        '| Service | Required create fields | Allowed updates / behavior |',
        '|---|---|---|',
        '| Assets | asset_name, asset_type, ip_address, owner_id, department, criticality | name, owner, department, criticality, status; optional operating_system/location on create; status defaults Active |',
        '| Threats | threat_type, threat_value, description, severity, source | description, severity, source; canonical IP/domain/hash; unique type/value; recent means latest five by date/id |',
        '| Vulnerabilities | title, description, cve_id, cvss_score, severity, asset_id, discovery_source | description, severity, CVSS, next status only; Open initially; resolution_date automatic on Resolved |',
        '| Playbooks | name, description, incident_category, severity | add actions with description/positive step_number and optional assigned_user_id; update action with next status and/or assigned_user_id; completion_date automatic |',
        '| Reports | no body | HTTP source stats only; admin/auditor; unavailable source returns 503 |',
        'Assets types: Server, Workstation, Database, Application, Network Device, Cloud Resource. Threat types follow common/enums.py exactly: IP Address, Domain, File Hash, Malware Signature, Other. File hashes accept 32/40/64 hex characters. CVEs match CVE-[four digits]-[at least four digits]. CVSS is a finite JSON number 0–10 with at most one decimal place, matching NUMERIC(3,1). Severity/criticality: Low, Medium, High, Critical.',
        'Vulnerability lifecycle: Open → Under Investigation → In Progress → Resolved → Closed. Action lifecycle: Pending → In Progress → Completed. Same-state, skipped and backward transitions are rejected. Assign an active user before action execution. Client-supplied history, completion dates and resolution dates are rejected.',
        'Search: Assets name/type/ip/department/criticality; Threats type/value/severity/source; Vulnerabilities cve/severity/asset/status. Text fragments match literally and case-insensitively; other values are exact. Multiple filters combine with AND. Empty searches, duplicate parameters, unknown fields and invalid values return 400. No matches return an empty list with 200.',
        'List/search responses use data: {count, assets|threats|vulnerabilities|playbooks: [...]}. Single resources return data containing their fields and id. A playbook detail includes actions sorted by step_number; action responses include execution_history. Reports contain report_type/generated_at/generated_by plus named totals/groupings. See the actual redacted request/response evidence for executed examples.'])
    (REPORT/'API_DOCUMENTATION.md').write_text('\n\n'.join(sections)+'\n',encoding='utf-8')
    (REPORT/'endpoint-inventory.json').write_text(json.dumps(inventories,indent=2)+'\n')


def evidence_pages(connection):
    style='''<style>body{font:16px Arial;margin:32px;color:#172033;background:#f7f9fc}h1{font-size:24px}h2{font-size:19px}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:white;border:1px solid #cbd5e1;border-radius:6px;padding:16px;font:14px Consolas;line-height:1.45}small{color:#475569}article{max-width:1120px}table{border-collapse:collapse;background:white}td,th{border:1px solid #cbd5e1;padding:8px;text-align:left}th{background:#dbeafe}</style>'''
    index=['<!doctype html><meta charset="utf-8">',style,'<h1>Actual execution evidence — Marc Bteich</h1>',
        '<p>Real Newman/Postman Runtime request/response records, captured locally. This is an execution report, not the Postman desktop interface. Secrets are redacted. Laboratory data is synthetic and was actually stored/executed.</p>']
    cards=[]
    for path in sorted(EVIDENCE.glob('newman-*.json')):
        result=json.loads(path.read_text())
        index.append(f'<h2>{html.escape(path.stem)}</h2><p>{result["requests"]["total"]} requests; {result["assertions"]["total"]} assertions; {len(result["failures"])} failures; {result["executed_at"]}</p>')
        if path.stem=='newman-combined':
            continue
        for number,record in enumerate(result['records'],1):
            if '/users/login' in record['url']:
                continue
            name=path.stem+f'-{number:02}'
            caption=f'Marc Bteich – {record["name"]} – {record["method"]} {urlparse(record["url"]).path}, HTTP {record["status"]}'
            content=['<!doctype html><meta charset="utf-8">',style,'<article>',
                '<h1>'+html.escape(caption)+'</h1>',
                '<small>Genuine Newman execution report • '+html.escape(result['executed_at'])+' • source: '+path.name+'</small>',
                '<p>'+html.escape(record['method']+' '+record['url'])+'</p>',
                '<h2>Actual request body (credentials redacted)</h2><pre>'+html.escape(json.dumps(record['request'],indent=2))+'</pre>',
                '<h2>Actual response — HTTP '+str(record['status'])+'</h2><pre>'+html.escape(json.dumps(record['response'],indent=2))+'</pre>',
                '<p>This screenshot displays stored results from an actual API execution. It is not a simulated API response or a Postman GUI screenshot.</p></article>']
            (EVIDENCE/(name+'.html')).write_text('\n'.join(content),encoding='utf-8')
            index.append(f'<p><a href="{name}.html">{html.escape(caption)}</a></p>')
            cards.append(dict(name=name,caption=caption,status=record['status'],method=record['method'],url=record['url']))
    # Inspect genuine local DB content; never expose password hashes.
    dbhtml=['<!doctype html><meta charset="utf-8">',style,'<h1>Marc Bteich – Database structure and stored records – actual local SQLite inspection</h1>']
    for (table,) in connection.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"):
        dbhtml.append('<h2>'+table+'</h2>')
        columns=connection.execute(f'PRAGMA table_info({table})').fetchall()
        dbhtml.append('<pre>'+html.escape(json.dumps(columns,indent=2))+'</pre>')
        names=[column[1] for column in columns if column[1]!='password_hash']
        rows=connection.execute('SELECT '+','.join(names)+' FROM '+table+' LIMIT 8').fetchall()
        dbhtml.append('<pre>'+html.escape(json.dumps([dict(zip(names,row)) for row in rows],indent=2,default=str))+'</pre>')
        dbhtml.append('<p>Foreign keys:</p><pre>'+html.escape(json.dumps(connection.execute(f'PRAGMA foreign_key_list({table})').fetchall(),indent=2))+'</pre>')
    (EVIDENCE/'database-inspection.html').write_text('\n'.join(dbhtml),encoding='utf-8')
    (EVIDENCE/'index.html').write_text('\n'.join(index),encoding='utf-8')
    (EVIDENCE/'screenshot-manifest.json').write_text(json.dumps(cards,indent=2)+'\n')
    return cards


def report_text(connection):
    paragraphs=['# Person B report draft'+ATTR,
        'Security Incident Detection and Response Management System — Sprint 1. Team: Elie Medawar and Marc Bteich. Prepared from the implemented backend and actual local tests on 20 October 2026. These are Marc’s assigned sections, to review and combine with Elie’s title page, contents, validation, authentication, GitHub sections and evidence. No claim is made that this draft has been personally reviewed or accepted by the instructor.',
        '## 3. Introduction'+ATTR,
        'The system provides a backend for recording and investigating security incidents, identifying affected organizational assets, maintaining threat indicators, tracking vulnerabilities and following response procedures. A central record of users, incidents, assets, audit events and notifications supports coordinated security work. This sprint implements REST endpoints, persistent storage, authentication, role-based permissions, validation and reproducible API tests; a frontend is outside the professor’s required scope.',
        'The objectives are to separate responsibilities into independent services, preserve relationships through database constraints, restrict operations to authorized roles, record security-relevant changes and produce statistics from service APIs. Marc owns Assets, Threat Intelligence, Vulnerabilities, Playbooks/Actions, three reports, the shared schema/sample data and documentation integration. Elie owns Users, Incidents, Audit Logs, Notifications, the shared foundation and the other three reports.',
        '## 4. System Architecture'+ATTR,
        'Nine Flask applications run independently on ports 5001–5009. They share one SQLite database and common Python infrastructure. Each service uses model classes for persistence, a service/validator layer for business rules and thin routes for authorization and responses. SQLAlchemy binds values rather than concatenating client input into SQL. SQLite development is permitted by the team plan; MySQL execution has not been verified.',
        'Public requests carry Bearer JWTs containing user_id, username, role and expiry. Shared role checks distinguish admin, analyst, user and auditor. Internal HTTP requests carry X-Service-Key. Assets verifies owners through Users; Incidents and Vulnerabilities verify assets through Assets; Playbooks verifies assigned users through Users. Every successful mutation requests an audit event. Critical threat/vulnerability creation identifies all admins through Users and requests notifications through Notifications.',
        'The Reports app registers both owners’ blueprints. It calls source /internal/stats endpoints through the existing get_service/ReportBuilder helpers and does not read source tables for report calculations. A missing required source yields 503. Logs and Notifications use the shared best-effort HTTP clients; normal writes can succeed if these recipients are unavailable. Playbook deletion is stricter: execution-history events must be acknowledged by Logs before actions can be cascaded away.',
        '![Architecture](diagrams/architecture.png)',
        '## 5. Implementation Details'+ATTR,
        '### Assets'+ATTR,
        'AssetValidator validates required strings, type, IP, owner ID, criticality and status. AssetService normalizes IPv4/IPv6, verifies the owner over HTTP and performs ORM writes and queries. Admins write; admins/auditors list; analysts can retrieve/search. Search filters combine with AND and escape LIKE wildcards. Incidents survive asset deletion with a null asset link; existing vulnerabilities block deletion with 409. The original Assets behavior is retained.',
        '### Threat Intelligence'+ATTR,
        'ThreatValidator canonicalizes IP addresses, domains (lowercase and no terminal dot), and hexadecimal hashes. Hash lengths identify MD5/SHA1/SHA256 formats; malware-signature/Other values must be nonempty strings. A unique constraint on threat_type/threat_value protects against duplicates including concurrent inserts. Admins write; analysts read/search; auditors read lists/details. The detailed Service 4 matrix is followed for auditor search, and the existing shared enum uses Other.',
        '### Vulnerability Management'+ATTR,
        'VulnerabilityValidator checks required fields, a CVE-YYYY-NNNN+ identifier, finite numeric CVSS 0–10 and a real asset. Only admins create/delete, while analysts may update. Status advances one step at a time through Open, Under Investigation, In Progress, Resolved and Closed. The server records resolution_date on Resolved. Invalid transitions are rejected before applying any field changes, so a rejected update cannot partially save a description.',
        '### Playbooks and Response Actions'+ATTR,
        'Playbooks organize response procedures by incident category and severity. Actions have unique positive step numbers within each playbook and are returned in order. Admins manage playbooks; admins/analysts add and progress actions. Pending advances to In Progress and then Completed; execution requires an active assigned user. completion_date is generated automatically. An append-only JSON execution_history field records actor, timestamp, prior/next status, assignment and completion. The JSON list is saved in the same transaction as the action. Before deletion, history is archived through Logs; a failed archive returns 503 and preserves the playbook. Deletion cascades to actions only after archival.',
        '### Reports and service communication'+ATTR,
        'The asset report returns total/critical assets and counts by type; the threat report returns total, counts by severity/type and the latest five records; the vulnerability report returns total/critical vulnerabilities and counts by severity/status. Source services calculate grouped counts with ORM aggregate queries. Admins and auditors may generate reports; analysts and regular users receive 403. Elie’s summary/incidents/activity routes are integrated unchanged.',
        '## 6. Database Design'+ATTR,
        'The schema contains exactly ten tables with integer primary keys. Foreign keys enforce ownership, reporting/assignment, affected assets, notes, notifications and response actions. Enum CHECK constraints restrict known roles, statuses, categories and severities; UNIQUE constraints protect usernames/emails, threat indicator pairs and playbook step numbers. CVSS and action-step bounds are checked by the database. IP/CVE formats and status transitions are validated by services because a membership CHECK cannot express an ordered lifecycle.',
        '![Entity relationships](diagrams/database.png)',
        'Users→Assets is one-to-many through owner_id; Users→Incidents uses reporter_id and optional assigned_analyst_id; Assets→Incidents and Assets→Vulnerabilities identify affected resources; Incidents→Investigation Notes is one-to-many; Users→Logs/Notifications identifies actors and recipients; Users→Playbooks identifies creators; Playbooks→Response Actions is ordered one-to-many; Users→Response Actions identifies assignees. Departments are fields rather than an extra table, preserving the agreed ten-table design.',
        'Deletion choices preserve useful records: incident asset/analyst links and log user links become null; log usernames remain as snapshots; incident notes and user notifications cascade with their parents; an asset owner and a vulnerable asset cannot be deleted while required references exist. Playbook actions cascade after their execution history is archived. Action-history schema migration adds one column to an existing table without dropping data. All ten ORM column definitions were compared with schema.sql.',
        '### Tables and field definitions'+ATTR]
    for (table,) in connection.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"):
        paragraphs.append('#### '+table+ATTR)
        columns=connection.execute(f'PRAGMA table_info({table})').fetchall()
        fk={r[3]:r for r in connection.execute(f'PRAGMA foreign_key_list({table})')}
        paragraphs += ['| Field | Type | Constraint / relationship |','|---|---|---|']
        for _,name,kind,required,default,pk in columns:
            constraints=['PK' if pk else 'NOT NULL' if required else 'nullable']
            if name in fk:
                f=fk[name];constraints.append(f'FK → {f[2]}.{f[4]}; delete {f[6]}')
            if default is not None:
                constraints.append('SQL default '+str(default))
            paragraphs.append(f'| {name} | {kind} | {"; ".join(constraints)} |')
    paragraphs.extend(['### Sample data and design decisions'+ATTR,
        'database.seed is additive and idempotent. It creates synthetic lab scenarios spanning all roles, account statuses, asset types/statuses, threat types/severities, incident/vulnerability lifecycles, notification types/read states and action statuses. Passwords are generated randomly, hashed with the existing User method and stored privately in .local/seed-credentials.json for local test access. No existing password is reset. Seeded lifecycle states/history are explicitly labelled scenario snapshots; HTTP tests separately demonstrate real transitions.',
        '## 8. API Documentation'+ATTR,
        'The merged endpoint reference is API_DOCUMENTATION.md. SIDRMS.postman_collection.json combines both students’ requests without altering Elie’s source collections or scripts; collection-level hooks are preserved as folder hooks. Marc’s individual collections exercise positive and negative cases and automatically create/clean relationship fixtures. Actual Newman results and redacted request/response records are in evidence/.',
        '## 11. References and AI-use Disclosure'+ATTR,
        'Professor’s Sprint1-Brief.docx; team PLAN.md and PROMPT_PERSON_B.md (kept outside Git). Flask documentation: https://flask.palletsprojects.com/en/stable/ . SQLAlchemy cascades: https://docs.sqlalchemy.org/en/20/orm/cascades.html . Python ipaddress: https://docs.python.org/3/library/ipaddress.html . Postman Newman: https://learning.postman.com/docs/reference/newman-cli/command-line-integration-with-newman/ . Accessed 20 October 2026.',
        'AI-use disclosure for Marc to review: OpenAI Codex assisted with interpreting the team contract, implementing Person B services and database scripts, writing/running automated and Newman tests, debugging integration, and preparing diagrams/documentation/evidence reports. Generated code and prose require Marc’s personal review, adaptation and understanding. Execution results come from actual local runs; synthetic seed data is labelled. No screenshots are represented as a Postman GUI session when they are browser captures of Newman execution reports.',
        '## Appendix: Evidence and submission review'+ATTR,
        'Review evidence/index.html, the PNG execution-report captures, database-inspection output and SCREENSHOT_CHECKLIST.md. The DOCX draft embeds selected genuine report screenshots. If the instructor requires the Postman desktop interface specifically, repeat those cases in Postman and replace/add GUI screenshots; Newman CLI results are clearly labelled. Combine these sections with Elie’s report and regenerate the common contents/page numbers. Both students must attend CREATES at least three days before the deadline, disclose AI assistance and personally upload the agreed final package after review.'])
    text='\n\n'.join(paragraphs)+'\n'
    (REPORT/'PERSON_B_REPORT.md').write_text(text,encoding='utf-8')


def capture_and_docx(cards):
    from playwright.sync_api import sync_playwright
    pngdir=REPORT/'screenshots/marc'
    pngdir.mkdir(parents=True,exist_ok=True)
    with sync_playwright() as playwright:
        browser=playwright.chromium.launch()
        page=browser.new_page(viewport={'width':1280,'height':900},device_scale_factor=1)
        for name in ('architecture','database'):
            page.goto((DIAGRAMS/(name+'.svg')).as_uri())
            page.locator('svg').screenshot(path=str(DIAGRAMS/(name+'.png')))
        # One genuine browser capture per executed B request. Large responses
        # are left intact in HTML/JSON; full-page captures preserve the output.
        for card in cards:
            page.goto((EVIDENCE/(card['name']+'.html')).as_uri())
            page.screenshot(path=str(pngdir/(card['name']+'.png')),full_page=True)
        page.goto((EVIDENCE/'database-inspection.html').as_uri())
        page.screenshot(path=str(pngdir/'database-inspection.png'),full_page=True)
        browser.close()
    from docx import Document
    from docx.shared import Inches
    from docx.shared import Pt
    def append_markdown(document,text):
        lines=text.splitlines()
        index=0
        code=False
        while index<len(lines):
            line=lines[index]
            index+=1
            if line.startswith('```'):
                code=not code
                continue
            if not line.strip():
                continue
            if code:
                run=document.add_paragraph().add_run(line)
                run.font.name='Consolas';run.font.size=Pt(9)
            elif line.startswith('|'):
                headers=[c.strip() for c in line.strip('|').split('|')]
                table=document.add_table(rows=1,cols=len(headers))
                table.style='Table Grid'
                for cell,value in zip(table.rows[0].cells,headers):
                    cell.text=value
                while index<len(lines) and (not lines[index].strip() or lines[index].startswith('|')):
                    if not lines[index].strip() or re.fullmatch(r'[| :\-]+',lines[index]):
                        index+=1
                        continue
                    values=[c.strip() for c in lines[index].strip('|').split('|')]
                    for cell,value in zip(table.add_row().cells,values):
                        cell.text=value
                    index+=1
            elif line.startswith('#'):
                count=len(line)-len(line.lstrip('#'))
                document.add_heading(line[count:].strip(),min(count,4))
            elif line.startswith('!['):
                file=re.search(r'\(([^)]+)\)',line).group(1)
                document.add_picture(str(REPORT/file),width=Inches(6.2))
            else:
                document.add_paragraph(line)
    document=Document()
    document.add_heading('Person B — Sprint 1 report draft (Marc Bteich)',0)
    append_markdown(document,(REPORT/'PERSON_B_REPORT.md').read_text(encoding='utf-8'))
    document.add_heading('Actual execution evidence (Marc Bteich)',1)
    # Embed successful writes, lifecycle steps, errors, searches and reports.
    # All remaining captures are supplied in the evidence folder.
    seen=set()
    for card in cards:
        service=urlparse(card['url']).port
        route=urlparse(card['url']).path
        key=(service,card['method'],card['status'], 'search' if '/search' in route else route if '/reports/' in route else '')
        if key in seen:
            continue
        seen.add(key)
        document.add_heading(card['caption'],2)
        document.add_paragraph('Browser capture of an actual Newman execution report; not a Postman GUI screenshot.')
        # Tall lists remain in the full-size screenshot/HTML rather than being
        # reduced to illegible text in the DOCX.
        from PIL import Image
        screenshot=pngdir/(card['name']+'.png')
        with Image.open(screenshot) as image:
            if image.height > 2300:
                document.add_paragraph('Full request/response capture: screenshots/marc/'+screenshot.name)
                continue
        document.add_picture(str(screenshot),width=Inches(6.0))
    document.add_heading('Merged API reference (Marc Bteich)',1)
    append_markdown(document,(REPORT/'API_DOCUMENTATION.md').read_text(encoding='utf-8'))
    document.save(REPORT/'Person_B_Report.docx')


def main():
    REPORT.mkdir(exist_ok=True)
    with closing(sqlite3.connect(ROOT/'database/sidrms.db')) as connection:
        diagrams(connection)
        route_documentation()
        cards=evidence_pages(connection)
        report_text(connection)
    capture_and_docx(cards)
    print('Generated report, merged API reference, diagrams and',len(cards),'genuine execution-report browser captures')


if __name__=='__main__':
    main()
